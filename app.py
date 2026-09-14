import json
import base64
from io import BytesIO
import mimetypes
import re
from datetime import date, datetime, time, timedelta
from pathlib import Path

import requests
from docx import Document
from pypdf import PdfReader
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
import streamlit as st
from streamlit.components.v1 import html


# Paste your OpenRouter API key between the quotation marks.
OPENROUTER_API_KEY = "ADD YOUR API KEY"
OPENROUTER_MODEL = "nvidia/nemotron-3.5-lightning:free"
OPENROUTER_VISION_MODEL = "google/gemini-2.0-flash-exp:free"
DATA_FILE = Path(__file__).with_name("smartassist_data.json")


st.set_page_config(page_title="SmartAssist", page_icon="🤖", layout="wide")


def load_data():
    """Load saved chats and tasks from the local JSON file."""
    if not DATA_FILE.exists():
        return {"chats": {"Chat 1": []}, "current_chat": "Chat 1", "tasks": []}

    try:
        data = json.loads(DATA_FILE.read_text(encoding="utf-8"))
        if "chats" not in data:
            data["chats"] = {"Chat 1": data.pop("messages", [])}
        if not data["chats"]:
            data["chats"] = {"Chat 1": []}
        data.setdefault("current_chat", next(iter(data["chats"])))
        data.setdefault("tasks", [])
        return data
    except (OSError, json.JSONDecodeError):
        return {"chats": {"Chat 1": []}, "current_chat": "Chat 1", "tasks": []}


def save_data():
    """Save chats and tasks so they survive a page refresh."""
    st.session_state.chats[st.session_state.current_chat] = st.session_state.messages
    data = {
        "chats": st.session_state.chats,
        "current_chat": st.session_state.current_chat,
        "tasks": st.session_state.tasks,
    }
    DATA_FILE.write_text(json.dumps(data, indent=2), encoding="utf-8")


if "chats" not in st.session_state:
    saved_data = load_data()
    st.session_state.chats = saved_data["chats"]
    st.session_state.current_chat = saved_data["current_chat"]
    st.session_state.messages = st.session_state.chats[st.session_state.current_chat]
    st.session_state.tasks = saved_data.get("tasks", [])


def ask_openrouter(messages, model=OPENROUTER_MODEL):
    """Send normalized messages to OpenRouter and return a clean reply."""
    if not OPENROUTER_API_KEY or OPENROUTER_API_KEY == "paste-your-openrouter-api-key-here":
        return "Please add your OpenRouter API key to OPENROUTER_API_KEY in app.py to use SmartAssist."

    headers = {
        "Authorization": f"Bearer {OPENROUTER_API_KEY}",
        "Content-Type": "application/json",
        "HTTP-Referer": "http://localhost:8502",
        "X-Title": "SmartAssist",
    }
    payload = {
        "model": model,
        "messages": messages,
        "reasoning": {"enabled": True},
    }
    for attempt in range(2):
        try:
            response = requests.post(
                "https://openrouter.ai/api/v1/chat/completions",
                headers=headers,
                json=payload,
                timeout=(15, 120),
            )
            response.raise_for_status()
            response_data = response.json()
            assistant_message = response_data["choices"][0]["message"]
            st.session_state.last_reasoning_details = assistant_message.get(
                "reasoning_details"
            )
            content = assistant_message.get("content", "")
            if isinstance(content, list):
                content = "\n".join(
                    item.get("text", "") for item in content if item.get("text")
                )
            return str(content).strip()
        except (requests.exceptions.Timeout, requests.exceptions.ConnectionError) as error:
            if attempt == 1:
                st.session_state.last_reasoning_details = None
                return (
                    "OpenRouter took too long to respond. The free model may be "
                    "busy. Please try again in a moment."
                )
        except requests.exceptions.RequestException as error:
            st.session_state.last_reasoning_details = None
            return f"I could not contact OpenRouter: {error}"
        except (KeyError, TypeError, ValueError) as error:
            st.session_state.last_reasoning_details = None
            return f"OpenRouter returned an unexpected response: {error}"


def extract_document_text(uploaded_file):
    """Extract structured text from PDF, Word, text, or Markdown files."""
    file_stream = BytesIO(uploaded_file.getvalue())
    file_name = uploaded_file.name.lower()
    if file_name.endswith(".pdf"):
        reader = PdfReader(file_stream)
        sections = [f"Title: {reader.metadata.title}" if reader.metadata and reader.metadata.title else ""]
        for page_number, page in enumerate(reader.pages, start=1):
            page_text = page.extract_text() or ""
            sections.append(f"--- Page {page_number} ---\n{page_text}")
        return "\n\n".join(section for section in sections if section)

    if file_name.endswith(".docx"):
        document = Document(file_stream)
        sections = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
        for table_number, table in enumerate(document.tables, start=1):
            rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
            sections.append(f"--- Table {table_number} ---\n" + "\n".join(rows))
        return "\n".join(sections)

    return uploaded_file.getvalue().decode("utf-8", errors="replace")


def create_file_message(uploaded_file, user_request):
    """Build the correct OpenRouter message for a text or image file."""
    file_name = uploaded_file.name.lower()
    if file_name.endswith(('.png', '.jpg', '.jpeg', '.webp')):
        mime_type = mimetypes.guess_type(uploaded_file.name)[0] or "image/png"
        encoded_file = base64.b64encode(uploaded_file.getvalue()).decode("ascii")
        return {
            "role": "user",
            "content": [
                {"type": "text", "text": user_request},
                {
                    "type": "image_url",
                    "image_url": {"url": f"data:{mime_type};base64,{encoded_file}"},
                },
            ],
        }, OPENROUTER_VISION_MODEL

    document_text = extract_document_text(uploaded_file)
    if not document_text.strip():
        raise ValueError("This file does not contain readable text.")
    return {
        "role": "user",
        "content": (
            f"The user attached {uploaded_file.name}. Use its contents to answer "
            f"the request.\n\nFile contents:\n{document_text}\n\n"
            f"User request: {user_request}"
        ),
    }, OPENROUTER_MODEL


def requested_download_format(message):
    lowered_message = message.lower()
    asks_for_file = any(
        phrase in lowered_message
        for phrase in ("make", "create", "generate", "give", "download")
    )
    if not asks_for_file:
        return None
    if "pdf" in lowered_message:
        return "pdf"
    if "word" in lowered_message or "docx" in lowered_message:
        return "docx"
    return None


def create_word_file(title, content):
    document = Document()
    document.add_heading(title, level=1)
    for paragraph in content.splitlines() or [content]:
        if paragraph.strip():
            document.add_paragraph(paragraph.strip())
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def create_pdf_file(title, content):
    """Create a valid PDF file in memory."""
    output = BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=letter,
        rightMargin=0.7 * inch,
        leftMargin=0.7 * inch,
        topMargin=0.7 * inch,
        bottomMargin=0.7 * inch,
    )
    styles = getSampleStyleSheet()
    story = [Paragraph(title, styles["Title"]), Spacer(1, 0.2 * inch)]
    for paragraph in content.splitlines() or [content]:
        if paragraph.strip():
            story.append(Paragraph(paragraph.strip(), styles["BodyText"]))
            story.append(Spacer(1, 0.08 * inch))
    document.build(story)
    return output.getvalue()


def extract_requested_task(message):
    match = re.search(
        r"(?:create|make|generate|write)\s+(?:a\s+)?(?:pdf|word|docx)\s+"
        r"(?:file\s+)?(?:for|of|containing|with)\s+(.+)",
        message,
        flags=re.IGNORECASE,
    )
    if match:
        return match.group(1).strip().rstrip(".!?")

    match = re.search(
        r"(?:create|make|generate|write)\s+(.+?)\s+(?:as|into)\s+a\s+"
        r"(?:pdf|word|docx)(?:\s+file)?",
        message,
        flags=re.IGNORECASE,
    )
    return match.group(1).strip().rstrip(".!?") if match else None


def switch_chat(chat_name):
    """Save the current chat and load another chat."""
    save_data()
    st.session_state.current_chat = chat_name
    st.session_state.messages = st.session_state.chats[chat_name]
    st.rerun()


def get_due_reminders():
    """Return active tasks whose reminder time has arrived."""
    now = datetime.now()
    reminders = []
    for task in st.session_state.tasks:
        reminder = task.get("reminder")
        if not task.get("completed") and reminder:
            try:
                if datetime.fromisoformat(reminder) <= now:
                    reminders.append(task)
            except ValueError:
                continue
    return reminders


def play_reminder_sound():
    """Play a short beep in the user's browser."""
    html(
        """
        <script>
        const audioContext = new (window.AudioContext || window.webkitAudioContext)();
        const oscillator = audioContext.createOscillator();
        const gain = audioContext.createGain();
        oscillator.frequency.value = 880;
        gain.gain.setValueAtTime(0.25, audioContext.currentTime);
        gain.gain.exponentialRampToValueAtTime(0.01, audioContext.currentTime + 0.7);
        oscillator.connect(gain);
        gain.connect(audioContext.destination);
        oscillator.start();
        oscillator.stop(audioContext.currentTime + 0.7);
        </script>
        """,
        height=0,
    )


@st.fragment(run_every="10s")
def reminder_alert():
    """Check for reminders regularly while the app is open."""
    if "alerted_reminders" not in st.session_state:
        st.session_state.alerted_reminders = set()

    due_reminders = get_due_reminders()
    new_reminders = []
    for task in due_reminders:
        reminder_id = f"{task['text']}|{task.get('reminder')}"
        if reminder_id not in st.session_state.alerted_reminders:
            st.session_state.alerted_reminders.add(reminder_id)
            new_reminders.append(task)

    if due_reminders:
        st.warning(
            "Reminder: " + ", ".join(task["text"] for task in due_reminders)
        )
    if new_reminders:
        play_reminder_sound()


def parse_reminder_time(time_text):
    """Convert an AM/PM time string into a time value."""
    try:
        return datetime.strptime(time_text.strip(), "%I:%M %p").time()
    except ValueError:
        return None


def update_reminder_from_message(message):
    """Update a matching task's reminder time from a chat command."""
    match = re.search(
        r"(?:update|change|edit|move|make)\s+(?:the\s+)?"
        r"(?:reminder\s+for\s+)?(.+?)\s+(?:to|at)\s+"
        r"(1[0-2]|[1-9])\s*(?::([0-5]\d))?\s*(am|pm)\b",
        message,
        flags=re.IGNORECASE,
    )
    if not match:
        return None

    task_name = match.group(1).strip(" .,!?'")
    task_name = re.sub(r"^(my|the)\s+", "", task_name, flags=re.IGNORECASE)
    hour = int(match.group(2))
    minute = int(match.group(3) or 0)
    meridiem = match.group(4).lower()
    if meridiem == "pm" and hour != 12:
        hour += 12
    if meridiem == "am" and hour == 12:
        hour = 0
    new_time = time(hour, minute)

    for task in st.session_state.tasks:
        if task_name.lower() in task["text"].lower():
            current_date = date.today()
            if task.get("reminder"):
                current_date = datetime.fromisoformat(task["reminder"]).date()
            task["reminder"] = datetime.combine(current_date, new_time).isoformat()
            save_data()
            return task["text"], new_time.strftime("%I:%M %p")

    return None


def add_tasks_from_message(message):
    """Create one or more tasks from reminder and to-do phrases in chat."""
    command = message.strip()
    lowered_command = command.lower()
    task_details = []

    reminder_match = re.search(
        r"(?:set\s+a\s+reminder\s+for|remind\s+me\s+about|remind\s+me\s+to)\s+(.+?)(?=,?\s+and\s+(?:add|also)|$)",
        command,
        flags=re.IGNORECASE,
    )
    if reminder_match:
        reminder_text = reminder_match.group(1).strip(" .,")
        reminder = datetime.combine(date.today(), time(9, 0))
        time_match = re.search(
            r"\b(1[0-2]|[1-9])\s*(?::([0-5]\d))?\s*(am|pm)\b",
            reminder_text,
            flags=re.IGNORECASE,
        )
        if time_match:
            hour = int(time_match.group(1))
            minute = int(time_match.group(2) or 0)
            if time_match.group(3).lower() == "pm" and hour != 12:
                hour += 12
            if time_match.group(3).lower() == "am" and hour == 12:
                hour = 0
            reminder = datetime.combine(date.today(), time(hour, minute))
        task_details.append((reminder_text, reminder))

    todo_match = re.search(
        r"add\s+['\"]?(.+?)['\"]?\s+to\s+my\s+to[- ]?do\s+list(?:\s+for\s+(today|tomorrow)(?:\s+morning)?)?",
        command,
        flags=re.IGNORECASE,
    )
    if todo_match:
        todo_text = todo_match.group(1).strip(" .,'\"")
        reminder = None
        schedule = (todo_match.group(2) or "").lower()
        if schedule == "today":
            reminder = datetime.combine(date.today(), time(9, 0))
        elif schedule == "tomorrow":
            reminder = datetime.combine(
                date.today() + timedelta(days=1), time(9, 0)
            )
        task_details.append((todo_text, reminder))

    if task_details:
        for task_text, reminder in task_details:
            st.session_state.tasks.append(
                {
                    "text": task_text,
                    "completed": False,
                    "reminder": reminder.isoformat() if reminder else None,
                }
            )
        save_data()
        return [task_text for task_text, _ in task_details]

    prefixes = (
        "remind me to ",
        "remind me ",
        "add reminder ",
        "add a reminder ",
        "add task ",
        "add a task ",
        "create task ",
        "create a task ",
        "add todo ",
        "add a todo ",
        "add to-do ",
        "add a to-do ",
    )

    task_text = ""
    for prefix in prefixes:
        if lowered_command.startswith(prefix):
            task_text = command[len(prefix):].strip()
            break

    if not task_text:
        return None

    reminder = None
    if task_text.lower().endswith(" tomorrow"):
        task_text = task_text[:-9].strip()
        reminder = datetime.combine(date.today() + timedelta(days=1), time(9, 0))
    elif task_text.lower().endswith(" today"):
        task_text = task_text[:-6].strip()
        reminder = datetime.combine(date.today(), time(9, 0))

    task_text = re.sub(r"^[,:-]\s*", "", task_text).strip()
    if not task_text:
        return None

    task_texts = [item.strip() for item in task_text.split(",") if item.strip()]
    for item in task_texts:
        st.session_state.tasks.append(
            {
                "text": item,
                "completed": False,
                "reminder": reminder.isoformat() if reminder else None,
            }
        )
    save_data()
    return task_texts


with st.sidebar:
    st.header("Chats")
    chat_names = list(st.session_state.chats)
    selected_chat = st.selectbox(
        "Open chat",
        chat_names,
        index=chat_names.index(st.session_state.current_chat),
    )
    if selected_chat != st.session_state.current_chat:
        switch_chat(selected_chat)

    new_chat_col, delete_chat_col = st.columns(2)
    with new_chat_col:
        if st.button("New Chat", use_container_width=True):
            new_chat_number = len(st.session_state.chats) + 1
            new_chat_name = f"Chat {new_chat_number}"
            while new_chat_name in st.session_state.chats:
                new_chat_number += 1
                new_chat_name = f"Chat {new_chat_number}"
            save_data()
            st.session_state.chats[new_chat_name] = []
            st.session_state.current_chat = new_chat_name
            st.session_state.messages = []
            save_data()
            st.rerun()

    with delete_chat_col:
        if st.button("Delete Chat", use_container_width=True):
            if len(st.session_state.chats) == 1:
                st.session_state.chats[st.session_state.current_chat] = []
                st.session_state.messages = []
            else:
                del st.session_state.chats[st.session_state.current_chat]
                st.session_state.current_chat = next(iter(st.session_state.chats))
                st.session_state.messages = st.session_state.chats[st.session_state.current_chat]
            save_data()
            st.rerun()

    st.header("Tasks & Reminders")

    with st.form("add_task_form", clear_on_submit=True):
        task_text = st.text_input("New task", placeholder="e.g. Call the dentist")
        add_reminder = st.checkbox("Set a reminder")
        reminder_date = st.date_input("Reminder date", value=date.today())
        reminder_time_text = st.text_input(
            "Reminder time (AM/PM)",
            value="09:00 AM",
            placeholder="e.g. 02:00 PM",
        )
        add_task = st.form_submit_button("Add Task", use_container_width=True)

    if add_task and task_text.strip():
        reminder = None
        if add_reminder:
            reminder_time = parse_reminder_time(reminder_time_text)
            if reminder_time is None:
                st.error("Use a time such as 02:00 PM or 09:30 AM.")
                st.stop()
            reminder = datetime.combine(reminder_date, reminder_time).isoformat()
        st.session_state.tasks.append(
            {
                "text": task_text.strip(),
                "completed": False,
                "reminder": reminder,
            }
        )
        save_data()
        st.rerun()

    active_tasks = [task for task in st.session_state.tasks if not task["completed"]]
    completed_count = len(st.session_state.tasks) - len(active_tasks)
    st.caption(f"{len(active_tasks)} active | {completed_count} completed")

    if not st.session_state.tasks:
        st.info("No tasks yet.")
    else:
        reminder_alert()

        for task_index, task in enumerate(st.session_state.tasks):
            is_completed = st.checkbox(
                task["text"],
                value=task["completed"],
                key=f"task_{task_index}",
            )
            if is_completed != task["completed"]:
                task["completed"] = is_completed
                save_data()
                st.rerun()

            st.markdown(f"**{task['text']}**")
            if task.get("reminder") and not task["completed"]:
                reminder_time_text = datetime.fromisoformat(
                    task["reminder"]
                ).strftime("%b %d at %I:%M %p")
                st.caption(f"Reminder time: {reminder_time_text}")
            elif not task["completed"]:
                st.caption("No reminder set")

            with st.form(f"rename_task_form_{task_index}"):
                renamed_task = st.text_input(
                    "Rename task",
                    value=task["text"],
                    key=f"rename_task_{task_index}",
                )
                if task.get("reminder"):
                    saved_reminder = datetime.fromisoformat(task["reminder"])
                    edited_date = st.date_input(
                        "Reminder date",
                        value=saved_reminder.date(),
                        key=f"edit_reminder_date_{task_index}",
                    )
                    edited_time_text = st.text_input(
                        "Reminder time (AM/PM)",
                        value=saved_reminder.strftime("%I:%M %p"),
                        key=f"edit_reminder_time_{task_index}",
                    )
                rename_col, delete_col = st.columns(2)
                with rename_col:
                    rename_task = st.form_submit_button(
                        "Rename", use_container_width=True
                    )
                with delete_col:
                    delete_task = st.form_submit_button(
                        "Delete", use_container_width=True
                    )
                update_reminder = st.form_submit_button(
                    "Update Reminder", use_container_width=True
                )

            if rename_task and renamed_task.strip():
                task["text"] = renamed_task.strip()
                save_data()
                st.rerun()

            if delete_task:
                st.session_state.tasks.pop(task_index)
                save_data()
                st.rerun()

            if update_reminder and task.get("reminder"):
                edited_time = parse_reminder_time(edited_time_text)
                if edited_time is None:
                    st.error("Use a time such as 08:55 PM or 09:30 AM.")
                else:
                    task["reminder"] = datetime.combine(
                        edited_date, edited_time
                    ).isoformat()
                    save_data()
                    st.rerun()


st.title("SmartAssist")
st.caption(f"Current chat: {st.session_state.current_chat}")

uploaded_file = st.file_uploader(
    "Attach a PDF, Word, text, Markdown, or image file to your chat",
    type=["pdf", "docx", "txt", "md", "png", "jpg", "jpeg", "webp"],
    help="Then ask SmartAssist to summarize it, explain it, or tell you what to do.",
)
if uploaded_file:
    st.caption(f"Attached: {uploaded_file.name}")

for message in st.session_state.messages:
    with st.chat_message(message["role"]):
        st.markdown(message["content"])

prompt = st.chat_input("What would you like help with?")
if prompt:
    st.session_state.messages.append({"role": "user", "content": prompt})
    st.session_state.last_reasoning_details = None
    with st.chat_message("user"):
        st.markdown(prompt)

    with st.chat_message("assistant"):
        updated_reminder = update_reminder_from_message(prompt)
        added_tasks = None if updated_reminder else add_tasks_from_message(prompt)
        download_format = requested_download_format(prompt)
        requested_task = extract_requested_task(prompt)
        if download_format and requested_task:
            answer = f"Created your {download_format.upper()} task file."
        elif updated_reminder:
            task_name, updated_time = updated_reminder
            answer = f"Updated the reminder for {task_name} to {updated_time}."
        elif added_tasks:
            task_count = len(added_tasks)
            task_word = "task" if task_count == 1 else "tasks"
            answer = f"Added {task_count} {task_word}: " + ", ".join(added_tasks)
        elif uploaded_file:
            try:
                file_message, file_model = create_file_message(uploaded_file, prompt)
            except Exception as error:
                answer = f"I could not read {uploaded_file.name}: {error}"
            else:
                request_messages = list(st.session_state.messages)
                request_messages[-1] = file_message
                with st.spinner("Reading the attached file..."):
                    answer = ask_openrouter(request_messages, model=file_model)
        else:
            answer = ask_openrouter(st.session_state.messages)
        st.markdown(answer)

        if download_format == "docx":
            st.download_button(
                "Download Word File",
                data=create_word_file(
                    "SmartAssist Task", requested_task or answer
                ),
                file_name="smartassist_response.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                key=f"download_docx_{st.session_state.current_chat}_{len(st.session_state.messages)}",
            )
        elif download_format == "pdf":
            st.download_button(
                "Download PDF File",
                data=create_pdf_file(
                    "SmartAssist Task", requested_task or answer
                ),
                file_name="smartassist_response.pdf",
                mime="application/pdf",
                key=f"download_pdf_{st.session_state.current_chat}_{len(st.session_state.messages)}",
            )

    assistant_message = {"role": "assistant", "content": answer}
    if st.session_state.last_reasoning_details:
        assistant_message["reasoning_details"] = (
            st.session_state.last_reasoning_details
        )
    st.session_state.messages.append(assistant_message)
    save_data()
